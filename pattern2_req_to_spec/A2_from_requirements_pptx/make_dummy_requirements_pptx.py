"""ダミーの「要件定義書（PowerPoint）」を作る。要件定義もパワポで持っている想定。
実在の製品・顧客とは無関係。numlib（バブルソート・掃き出し法）の要件を R-xx / N-xx の ID 付きで書く。"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
OUT = Path(__file__).parent / "input"; OUT.mkdir(exist_ok=True)
prs = Presentation()
def title(t, sub):
    s = prs.slides.add_slide(prs.slide_layouts[0]); s.shapes.title.text = t; s.placeholders[1].text = sub
def bullets(t, lines):
    s = prs.slides.add_slide(prs.slide_layouts[1]); s.shapes.title.text = t
    tf = s.placeholders[1].text_frame; tf.clear()
    for i, (lvl, txt) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.text = txt; p.level = lvl; p.font.size = Pt(18)
def table(t, header, rows, col_w):
    s = prs.slides.add_slide(prs.slide_layouts[5]); s.shapes.title.text = t
    tb = s.shapes.add_table(len(rows) + 1, len(header), Inches(0.4), Inches(1.4), Inches(9.2), Inches(0.4)).table
    for c, h in enumerate(header): tb.cell(0, c).text = h
    for r, row in enumerate(rows, 1):
        for c, v in enumerate(row): tb.cell(r, c).text = v
    for r in range(len(rows) + 1):
        for c in range(len(header)):
            for p in tb.cell(r, c).text_frame.paragraphs: p.font.size = Pt(12)
    for c, w in enumerate(col_w): tb.columns[c].width = Inches(w)
title("数値処理ライブラリ numlib 要件定義書", "版数 1.0 ／ 2026-09-25 ／ 制御ソフト第二課\n※ AI検証用のダミー資料。実在の製品・顧客とは無関係")
bullets("1. 背景と目的", [
    (0, "背景：制御ソフト各機能で「整列」と「連立一次方程式の求解」が個別に実装され、品質と保守性にばらつきがある"),
    (0, "目的：共通ライブラリ numlib として一本化し、単体試験で品質を担保する"),
    (0, "期待効果"), (1, "重複コードの削減"), (1, "試験の標準化（試験観点を仕様書に含める）"), (1, "新人でも安全に使える API"),
])
bullets("2. 適用範囲と前提", [
    (0, "対象：組込み制御ソフト（C99）。PC 向けツールは対象外"),
    (0, "前提"), (1, "ヒープ（malloc/free）は使えない"), (1, "標準入出力は製品コードでは使わない"),
    (1, "シングルスレッドで呼ばれる（排他は呼び出し側）"), (1, "浮動小数点は double を使用可"), (1, "依存ヘッダは最小限（stddef.h / stdint.h 程度）"),
])
table("3. 機能要件", ["ID", "要件", "備考"], [
    ["R-01", "32bit 符号付き整数の配列を昇順に整列できること", "同じ値の相対順序を保つ（安定）"],
    ["R-02", "要素数 0 または 1 の配列でも正常終了すること", "何もしないで成功を返す"],
    ["R-03", "NULL ポインタや不正な長さを渡してもクラッシュせず、エラーを返すこと", "全関数共通"],
    ["R-04", "連立一次方程式 Ax = b（n ≦ 16）の解 x を求められること", "A は n×n、b・x は長さ n"],
    ["R-05", "求解では部分ピボット選択を行うこと", "対角に 0 があっても解けること"],
    ["R-06", "特異行列（ピボットの絶対値が 1e-12 未満）を検出し、エラーを返すこと", "解を書き込まない"],
    ["R-07", "条件数 1e6 以下の行列で、解の相対誤差が 1e-9 以内であること", "受け入れ試験で確認"],
], [0.8, 5.6, 2.8])
table("4. 非機能要件", ["ID", "要件", "備考"], [
    ["N-01", "動的メモリ確保を行わないこと", "静的・スタックのみ"],
    ["N-02", "製品コードで標準入出力を使わないこと", "試験コードは可"],
    ["N-03", "-Wall -Wextra -Wpedantic で警告ゼロ", "C99"],
    ["N-04", "計算量：整列は O(n²) 以内、求解は O(n³) 以内", ""],
    ["N-05", "スタック使用量は n ≦ 16 を前提とした定数以内", "再帰禁止"],
    ["N-06", "スレッドセーフは不要", "呼び出し側で排他"],
], [0.8, 5.6, 2.8])
bullets("5. インターフェース方針", [
    (0, "エラーは戻り値で返す：0 = 成功、負の値 = 失敗。失敗の種別（引数不正／特異行列／サイズ超過）を区別できること"),
    (0, "行列は行優先（row-major）の 1 次元配列 double[n*n] で渡す（ポインタ配列 double** は使わない）"),
    (0, "求解関数は入力 A・b を作業領域として破壊してよい（呼び出し側が必要ならコピーする）"),
    (0, "関数名・定数名は接頭辞 numlib_ / NUMLIB_ で統一する"),
])
table("6. 受け入れ基準", ["項目", "基準"], [
    ["単体試験", "R-01〜R-07 のそれぞれを最低 1 件の試験で検証していること"],
    ["カバレッジ", "行カバレッジ 90% 以上（gcov で計測）"],
    ["仕様書", "各関数の引数・戻り値・事後条件・失敗時の状態、および試験観点を含むこと"],
    ["トレース", "要件 ID と関数・試験観点の対応が表で追えること"],
    ["納期", "2026-10-31（仕様書レビュー 10-10、実装完了 10-24）"],
], [2.0, 7.2])
table("7. 用語・改訂履歴", ["項目", "内容"], [
    ["部分ピボット選択", "各列で絶対値最大の要素を持つ行を選び、行を入れ替えてから消去する方法"],
    ["特異行列", "逆行列が存在しない（解が一意に定まらない）行列"],
    ["改訂 1.0", "2026-09-25 初版（制御ソフト第二課）"],
], [2.4, 6.8])
p = OUT / "numlib_要件定義書_v1.0.pptx"; prs.save(p); print("wrote", p.name, p.stat().st_size, "bytes")
